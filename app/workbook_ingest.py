"""월 마감 실적 워크북 인입 모듈.

**최종 백데이터 형식(유일하게 지원하는 형식)**: "DC율"·"프로모션 기준"·"DATA" 3개
시트로 구성된 워크북(`.docs/17_실데이터_워크북_반영_계획서.md`·`19_최종백데이터_전환_계획서.md`
참고). 예전에 있던 총괄표/백데이터/제품표 구조 경로는 완전히 걷어냈다 — 앞으로 업로드되는
자료는 전부 이 형식이라는 사용자 확인에 따른 것이다.

**핵심 설계 원칙**: 프로모션 조건이 매달 바뀌어도 코드를 고치지 않고 "프로모션 기준" 시트
값만 바꿔서 반영되어야 한다. 그래서 산식은 프로모션명(예: "대형"·"특화")이라는 자유
텍스트가 아니라, **"산식유형"이라는 고정 어휘 컬럼**(무상증정/품목별고정단가/구간별단가-
절대수량/연동형)으로 분기한다. 연동형은 **"트리거제품군"** 컬럼으로 어떤 제품군의
판매수량을 볼지 정한다(더 이상 "AGM"으로 고정돼 있지 않다).

- 대리점별 고정 DC(매출DC+성장DC)는 **PRD 6.3.3 정의 그대로 "기준가 × 판매수량 × DC율"**을
  (대리점,제품) 단위로 계산한다. "DC율 미적용" 제품은 매출DC·성장DC 둘 다 0원이다.
- 프로모션 지원금액은 "프로모션 기준" 시트의 산식유형별로 이 모듈이 직접 계산한다.
- **영업이익·매출원가는 DATA 시트 값을 그대로 쓴다**(재계산하지 않음) — 실제 회계 결과이기
  때문이다.

한때 매출액-매출원가-영업이익로 역산한 값과 계산값을 나란히 보여주는 "역산 검증"을
제공했으나, 실제 "26년 8월 DC 총괄표.xlsx"(공식 집계 자료)와 직접 대조해보니 우리
계산값 자체는 거의 정확했고(대형 대리점 6곳 검증, 원본 대비 오차 0~0.1% 수준), 그
방식(매출-원가-영업이익)은 부가세 등 다른 고정비가 섞여 있어 신뢰도가 낮다는 게
확인돼 MVP에서 뺐다(사용자 확인). 검증이 다시 필요해지면 이런 공식 총괄표류 자료와
직접 대조하는 편이 훨씬 정확하다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

import openpyxl
import pandas as pd


class ColumnNotFoundError(ValueError):
    pass


@dataclass
class IngestResult:
    sales_like: pd.DataFrame
    applied_support: pd.DataFrame
    product_group_support: pd.DataFrame
    promotion_notice: list[str] = field(default_factory=list)
    threshold_proximity: pd.DataFrame = field(default_factory=pd.DataFrame)
    promotion_insight: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _locate_header_row(ws, must_include: list[str], max_scan_rows: int = 10):
    """상단 `max_scan_rows`행 안에서 `must_include`에 있는 텍스트를 모두 포함하는
    행을 찾아 (행 번호, {셀 텍스트: 열 번호}) 를 돌려준다. 못 찾으면 (None, None)."""
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        row_texts: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                row_texts[str(v).strip()] = c
        if all(any(k in text for text in row_texts) for k in must_include):
            return r, row_texts
    return None, None


def _find_매출액_subcolumns(ws, header_row: int) -> dict[str, int] | None:
    """"매출액" 병합 슈퍼헤더 아래의 실제 열 라벨은 "26년 8월"·"26년 7월"·"25년 8월"처럼
    매달 바뀌는 상대 연월 문자열이라, "전월매출"·"전년동월매출" 같은 고정 텍스트로는
    찾을 수 없다. 대신 "매출액" 병합 범위 안에서 **정해진 열 순서**(이번달매출 →
    전월매출 → 전월대비증감율 → 전년동월매출 → 전년동월대비증감율)로 찾는다 — 이 5개
    열은 팀이 매달 값만 새로 채우고 순서는 바꾸지 않는 고정 양식이라는 전제다.
    "매출액" 슈퍼헤더가 아예 없는 워크북(예전 DC율 시트)에서는 None을 돌려주고, 호출부는
    전월매출/전년동월매출 없이 그레이스풀하게 진행한다."""
    for merged in ws.merged_cells.ranges:
        if merged.min_row != header_row:
            continue
        top_left = ws.cell(row=merged.min_row, column=merged.min_col).value
        if top_left and "매출액" in str(top_left).strip():
            width = merged.max_col - merged.min_col + 1
            if width < 5:
                return None
            start = merged.min_col
            return {
                "이번달매출": start,
                "전월매출": start + 1,
                "전월대비증감율": start + 2,
                "전년동월매출": start + 3,
                "전년동월대비증감율": start + 4,
            }
    return None


def extract_DC율(wb) -> pd.DataFrame:
    """대리점별 매출DC·성장DC 비율과(있으면) 전월매출·전년동월매출."""
    ws = wb["DC율"]
    header_row, header_texts = _locate_header_row(ws, ["거래처코드"])
    if header_row is None:
        raise ColumnNotFoundError('"DC율" 시트에서 "거래처코드" 헤더를 찾지 못했습니다.')

    def _find(*keywords: str) -> int | None:
        return next((col for text, col in header_texts.items() if all(k in text for k in keywords)), None)

    col_store, col_name = _find("거래처코드"), _find("거래처명")
    col_sale_dc, col_growth_dc = _find("매출DC"), _find("성장DC")
    if None in (col_store, col_sale_dc, col_growth_dc):
        raise ColumnNotFoundError('"DC율" 시트에 필요한 컬럼(거래처코드/매출DC/성장DC)이 없습니다.')

    매출액_cols = _find_매출액_subcolumns(ws, header_row)
    col_prev_revenue = 매출액_cols["전월매출"] if 매출액_cols else None
    col_prev_year_revenue = 매출액_cols["전년동월매출"] if 매출액_cols else None

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        store = ws.cell(row=r, column=col_store).value
        if store is None:
            continue
        rows.append(
            {
                "대리점코드": str(store).strip(),
                "대리점명": ws.cell(row=r, column=col_name).value if col_name else None,
                "매출DC율": ws.cell(row=r, column=col_sale_dc).value or 0,
                "성장DC율": ws.cell(row=r, column=col_growth_dc).value or 0,
                "전월매출": ws.cell(row=r, column=col_prev_revenue).value if col_prev_revenue else None,
                "전년동월매출": ws.cell(row=r, column=col_prev_year_revenue).value if col_prev_year_revenue else None,
            }
        )
    return pd.DataFrame(rows)


def _extract_프로모션기준(ws) -> tuple[pd.DataFrame, list[dict], pd.DataFrame]:
    """"프로모션 기준" 시트를 읽는다. 같은 제품코드가 "특화"+"대형"처럼 두 번 나오면
    (예: GB 73010 — 회색 음영 "특화+대형 동시 적용" 사례) 제품 정보는 한 번만, 규칙은
    각각 별도로 쌓는다. 시트 오른쪽의 "AGM Level 별 달성 수량"(대리점별 Lv.1~3 문턱값)
    표도 같은 헤더 행에서 함께 읽는다.

    **"산식유형" 컬럼이 반드시 있어야 한다** — 프로모션명(자유 텍스트, 화면 표시용)이
    아니라 이 컬럼의 고정 어휘(무상증정/품목별고정단가/구간별단가-절대수량/연동형)로
    산식을 분기해야, 프로모션명이 매달 바뀌어도 코드 수정 없이 계산이 유지된다."""
    header_row, header_texts = _locate_header_row(ws, ["제품코드", "기준가"])
    if header_row is None:
        raise ColumnNotFoundError('"프로모션 기준" 시트에서 "제품코드"/"기준가" 헤더를 찾지 못했습니다.')

    def _find(*keywords: str) -> int | None:
        return next((col for text, col in header_texts.items() if all(k in text for k in keywords)), None)

    col_code, col_name = _find("제품코드"), _find("제품명")
    col_group1, col_group2 = _find("제품구분1"), _find("제품구분2")
    col_price, col_note = _find("기준가"), _find("비고")
    col_promo_name = _find("프로모션명")
    col_formula_type = _find("산식유형")
    col_trigger_group = _find("트리거제품군")
    col_trigger_amount_start = _find("트리거제품군", "지원금액")
    col_cond_start = _find("지원", "조건")
    col_amount_start = _find("지원", "금액")
    if None in (col_code, col_price, col_promo_name, col_cond_start, col_amount_start):
        raise ColumnNotFoundError('"프로모션 기준" 시트에 필요한 컬럼이 없습니다.')
    if col_formula_type is None:
        raise ColumnNotFoundError(
            '"프로모션 기준" 시트에 "산식유형" 컬럼이 없습니다 — 프로모션명(자유 텍스트)이 아니라 '
            "산식유형(무상증정/품목별고정단가/구간별단가-절대수량/연동형)을 명시해야 계산할 수 있습니다."
        )

    product_rows: dict[str, dict] = {}
    rules: list[dict] = []
    for r in range(header_row + 1, ws.max_row + 1):
        code = ws.cell(row=r, column=col_code).value
        if code is None:
            continue
        code = str(code).strip()
        if code not in product_rows:
            note = ws.cell(row=r, column=col_note).value if col_note else None
            product_rows[code] = {
                "제품코드": code,
                "제품명": ws.cell(row=r, column=col_name).value if col_name else None,
                "제품구분1": ws.cell(row=r, column=col_group1).value if col_group1 else None,
                "제품구분2": ws.cell(row=r, column=col_group2).value if col_group2 else None,
                "기준가": ws.cell(row=r, column=col_price).value or 0,
                "DC미적용": bool(note and "미적용" in str(note)),
            }
        formula_type = ws.cell(row=r, column=col_formula_type).value
        if not formula_type:
            continue
        promo_name = ws.cell(row=r, column=col_promo_name).value
        conds = [ws.cell(row=r, column=col_cond_start + i).value for i in range(3)]
        amounts = [ws.cell(row=r, column=col_amount_start + i).value for i in range(3)]
        trigger_group = ws.cell(row=r, column=col_trigger_group).value if col_trigger_group else None
        trigger_amounts = (
            [ws.cell(row=r, column=col_trigger_amount_start + i).value for i in range(3)]
            if col_trigger_amount_start
            else [None, None, None]
        )
        rules.append(
            {
                "제품코드": code,
                "프로모션명": str(promo_name).strip() if promo_name else str(formula_type).strip(),
                "산식유형": str(formula_type).strip(),
                "트리거제품군": str(trigger_group).strip() if trigger_group else None,
                "conds": conds,
                "amounts": amounts,
                "트리거지원금액": trigger_amounts,
            }
        )

    제품마스터 = pd.DataFrame(product_rows.values())

    col_agm_code, col_agm_name = header_texts.get("코드"), header_texts.get("대리점")
    col_lv1, col_lv2, col_lv3 = header_texts.get("AGM Lv.1"), header_texts.get("AGM Lv.2"), header_texts.get("AGM Lv.3")
    agm_rows = []
    if col_agm_code and col_lv1 and col_lv2 and col_lv3:
        for r in range(header_row + 1, ws.max_row + 1):
            store = ws.cell(row=r, column=col_agm_code).value
            if store is None:
                continue
            agm_rows.append(
                {
                    "대리점코드": str(store).strip(),
                    "Lv1": ws.cell(row=r, column=col_lv1).value or 0,
                    "Lv2": ws.cell(row=r, column=col_lv2).value or 0,
                    "Lv3": ws.cell(row=r, column=col_lv3).value or 0,
                }
            )
    agm_lv = pd.DataFrame(agm_rows)

    return 제품마스터, rules, agm_lv


def _infer_규칙(rule: dict) -> dict:
    """제품별 프로모션 규칙 한 줄을 "산식유형" 컬럼 값으로 해석한다(더 이상
    프로모션명 문자열을 추측하지 않는다 — 프로모션명이 바뀌어도 이 컬럼만 정확하면
    계산이 그대로 유지된다)."""
    유형 = rule["산식유형"]
    conds, amounts = rule["conds"], rule["amounts"]
    if 유형 == "품목별고정단가":
        return {"유형": 유형, "단가": amounts[0] or 0}
    if 유형 == "구간별단가-절대수량":
        tiers = []
        for cond, amt in zip(conds, amounts):
            if cond is None or amt is None:
                continue
            m = re.search(r"(\d+)", str(cond))
            if m:
                tiers.append((int(m.group(1)), amt))
        tiers.sort(key=lambda t: t[0])
        return {"유형": 유형, "tiers": tiers}
    if 유형 == "연동형":
        return {
            "유형": 유형,
            "tier_rates": [a or 0 for a in amounts],
            "트리거제품군": rule.get("트리거제품군"),
            # 트리거제품군(예: AGM) 자체도 Lv 달성 시 대당 지원금액을 별도로 받는다 —
            # 대상 품목(GB 등)에 붙는 tier_rates와는 다른 단가표다.
            "트리거_tier_rates": [a or 0 for a in rule.get("트리거지원금액") or [0, 0, 0]],
        }
    if 유형 == "무상증정":
        m = re.search(r"(\d+)", str(conds[0])) if conds and conds[0] else None
        n = int(m.group(1)) if m else int(amounts[0] or 1)
        return {"유형": 유형, "n": n}
    return {"유형": "미정의"}


def extract_DATA(wb, 제품코드_집합: set[str]) -> tuple[pd.DataFrame, list[str]]:
    """"DATA" 시트를 (대리점, 제품) 단위로 집계한다(같은 조합이 여러 줄이면 합산 —
    한 달 안에 여러 건의 전표가 있을 수 있다). 프로모션 기준 제품마스터에 없는
    코드(예: 이륜용 제품)는 제외하고 그 코드 목록을 돌려준다.

    ERP에서 내려받는 원본이라 컬럼명이 내보낼 때마다 바뀔 수 있다(예: "제품코드"가
    "상품"으로, "대리점명"이 "고객"으로 바뀐 적이 있었다) — 그래서 컬럼마다 후보
    키워드를 여러 개 순서대로 시도한다. "매출원가"처럼 비슷한 이름의 컬럼이 여럿이면
    (예: "매출원가(A)Tot"·"매출원가(A)") 더 구체적인 키워드(Tot 포함)를 우선한다."""
    ws = wb["DATA"]
    header_row, header_texts = _locate_header_row(ws, ["거래처코드", "매출액"])
    if header_row is None:
        raise ColumnNotFoundError('"DATA" 시트에서 필요한 헤더를 찾지 못했습니다.')

    def _find(*keyword_groups: tuple[str, ...]) -> int | None:
        for keywords in keyword_groups:
            match = next((col for text, col in header_texts.items() if all(k in text for k in keywords)), None)
            if match is not None:
                return match
        return None

    col_store = _find(("거래처코드",))
    col_store_name = _find(("대리점명",), ("고객",))
    col_code = _find(("제품코드",), ("상품",))
    col_revenue = _find(("매출액",))
    col_qty = _find(("매출수량",))
    col_profit = _find(("영업이익",))
    col_cost = _find(("매출원가", "Tot"), ("매출원가",))
    required = [col_store, col_code, col_revenue, col_qty, col_profit, col_cost]
    if any(c is None for c in required):
        raise ColumnNotFoundError('"DATA" 시트에 필요한 컬럼(거래처코드/제품코드 또는 상품/매출액/매출수량/영업이익/매출원가)이 없습니다.')

    rows = []
    excluded_codes: set[str] = set()
    for r in range(header_row + 1, ws.max_row + 1):
        store = ws.cell(row=r, column=col_store).value
        code = ws.cell(row=r, column=col_code).value
        if store is None or code is None:
            continue
        code = str(code).strip()
        if code not in 제품코드_집합:
            excluded_codes.add(code)
            continue
        rows.append(
            {
                "대리점코드": str(store).strip(),
                "대리점명": ws.cell(row=r, column=col_store_name).value if col_store_name else None,
                "제품코드": code,
                "매출액": ws.cell(row=r, column=col_revenue).value or 0,
                "판매수량": ws.cell(row=r, column=col_qty).value or 0,
                "영업이익": ws.cell(row=r, column=col_profit).value or 0,
                "매출원가": ws.cell(row=r, column=col_cost).value or 0,
            }
        )
    df = pd.DataFrame(rows)
    if df.empty:
        return df, sorted(excluded_codes)
    agg = df.groupby(["대리점코드", "제품코드"], as_index=False).agg(
        대리점명=("대리점명", "first"),
        매출액=("매출액", "sum"),
        판매수량=("판매수량", "sum"),
        영업이익=("영업이익", "sum"),
        매출원가=("매출원가", "sum"),
    )
    return agg, sorted(excluded_codes)


def _tier_lookup(qty: float, tiers: list[tuple[float, float]]) -> tuple[float, float | None, float | None]:
    """`tiers`(문턱값 오름차순 [(수량, 단가), ...])에서 `qty`가 도달한 가장 높은
    구간의 단가와, 아직 못 넘었으면 다음 구간의 문턱값·단가를 함께 돌려준다.
    (현재단가, 다음문턱값 또는 None, 다음단가 또는 None) — 구간별단가·연동형(AGM Lv)
    둘 다 "수량으로 구간을 찾는다"는 점이 같아서 이 함수 하나로 같이 쓴다."""
    current_rate = 0.0
    for i, (threshold, amount) in enumerate(tiers):
        if qty >= threshold:
            current_rate = amount
        else:
            return current_rate, threshold, amount
    return current_rate, None, None


def _support_rows(store_df: pd.DataFrame, 기준연월: str, 지원유형: str, amounts: pd.Series) -> pd.DataFrame:
    """지원금액이 0이 아닌 대리점만 골라 applied_support 형태(기준연월/대리점코드/
    대리점명/제품코드/지원유형/지원단가/지원금액)의 롱포맷 DataFrame으로 만든다.
    그레인이 대리점 단위라 제품코드·지원단가는 의미가 없어 항상 비운다."""
    mask = amounts != 0
    columns = ["기준연월", "대리점코드", "대리점명", "제품코드", "지원유형", "지원단가", "지원금액"]
    if not mask.any():
        return pd.DataFrame(columns=columns)
    return pd.DataFrame(
        {
            "기준연월": 기준연월,
            "대리점코드": store_df.loc[mask, "대리점코드"],
            "대리점명": store_df.loc[mask, "대리점명"],
            "제품코드": None,
            "지원유형": 지원유형,
            "지원단가": pd.NA,
            "지원금액": amounts.loc[mask],
        }
    )


def _build_promotion_notice(rules: list[dict]) -> list[str]:
    """프로모션명별로 묶어 "총 n개 품목" 수준으로 요약한다(제품코드를 나열하지 않음).

    같은 프로모션명 안에서도 제품코드별로 실제 지원 단가가 다를 수 있다 — 예를 들어
    "대형"은 대부분 품목이 대당 3,000원이지만, "특화"와 동시 적용되는 4개 품목
    (GB73010/73011/250L/250R)만 대당 4,000원이다. 대표 한 행(`items[0]`)만 보고
    요약하면 이런 변형이 누락되므로, 그룹 안의 **모든** 행을 계산해 서로 다른 값이
    나오면 각각 몇 개 품목인지와 함께 보여준다."""
    grouped: dict[str, list[dict]] = {}
    for rule in rules:
        grouped.setdefault(rule["프로모션명"], []).append(rule)

    lines: list[str] = []
    for name, items in grouped.items():
        lines.append(f"▶ {name} ({len(items)}개 품목)")
        infos = [_infer_규칙(item) for item in items]
        유형 = infos[0]["유형"]

        if 유형 == "품목별고정단가":
            단가별_품목수: dict[float, int] = {}
            for info in infos:
                단가별_품목수[info["단가"]] = 단가별_품목수.get(info["단가"], 0) + 1
            if len(단가별_품목수) == 1:
                (단가,) = 단가별_품목수
                lines.append(f"  대당 {단가:,.0f}원 지원 (수량 문턱값 없음)")
            else:
                for 단가, 개수 in sorted(단가별_품목수.items()):
                    lines.append(f"  대당 {단가:,.0f}원 지원 (수량 문턱값 없음) — {개수}개 품목")
        elif 유형 == "구간별단가-절대수량":
            unique_tiers = {tuple(info["tiers"]) for info in infos}
            for tiers in sorted(unique_tiers):
                tiers_text = " / ".join(f"{threshold}대 이상 {amount:,.0f}원" for threshold, amount in tiers)
                개수 = sum(1 for info in infos if tuple(info["tiers"]) == tiers)
                suffix = f" — {개수}개 품목" if len(unique_tiers) > 1 else ""
                lines.append(f"  판매수량 구간별 대당 지원: {tiers_text}{suffix}")
        elif 유형 == "연동형":
            info = infos[0]
            rates_text = " / ".join(f"Lv.{i + 1} {rate:,.0f}원" for i, rate in enumerate(info["tier_rates"]))
            트리거 = info.get("트리거제품군") or "?"
            lines.append(f"  대리점별 {트리거} 판매수량이 Lv 문턱값을 넘으면 지정 품목에 대당 지원: {rates_text}(대리점마다 Lv 문턱값이 다름)")
            트리거_rates = info.get("트리거_tier_rates") or [0, 0, 0]
            if any(트리거_rates):
                트리거_rates_text = " / ".join(f"Lv.{i + 1} {rate:,.0f}원" for i, rate in enumerate(트리거_rates))
                lines.append(f"  {트리거} 자체도 Lv 달성 시 대당 지원(트리거 제품군 자체지원): {트리거_rates_text}")
        elif 유형 == "무상증정":
            lines.append(f"  판매수량 {infos[0]['n']}대당 1대분(기준가) 지원")
    return lines


def _build_threshold_proximity(
    data: pd.DataFrame, rules: list[dict], agm_lv: pd.DataFrame, trigger_qty_by_store: dict[str, pd.Series]
) -> pd.DataFrame:
    """구간별단가(절대수량)·연동형처럼 "문턱값을 넘었는지"로 지원이 갈리는 규칙에
    대해, 다음 구간까지 얼마나 남았는지 보여준다 — 이 프로모션이 추가 구매를 실제로
    유도하고 있는지, 다음 구간까지 조금만 더 팔면 되는 대리점이 있는지 직접 보여주는
    지표다(`.docs/16_교차검증_보완_계획서.md`·`19_최종백데이터_전환_계획서.md`)."""
    rows: list[dict] = []

    # 구간별단가(절대수량)는 프로모션명 그룹에 속한 "모든" 제품코드의 판매수량을
    # 대리점 단위로 합산한 값으로 구간을 판정한다(개별 제품코드 단위로 판정하면
    # 틀림 — CLAUDE.md 핵심 도메인 규칙, `build_sales_like_table`의 실제 계산과
    # 동일한 집계 기준을 써야 근접도 표시도 일관된다).
    tiered_group_codes: dict[str, list[str]] = {}
    for rule in rules:
        if _infer_규칙(rule)["유형"] == "구간별단가-절대수량":
            tiered_group_codes.setdefault(rule["프로모션명"], []).append(rule["제품코드"])

    store_names = data.drop_duplicates("대리점코드").set_index("대리점코드")["대리점명"]
    for label, codes in tiered_group_codes.items():
        info = next(_infer_규칙(r) for r in rules if r["프로모션명"] == label)
        if not info["tiers"]:
            continue
        mask = data["제품코드"].isin(codes)
        qty_by_store = data.loc[mask].groupby("대리점코드")["판매수량"].sum()
        for store, qty in qty_by_store.items():
            현재단가, 다음문턱값, 다음단가 = _tier_lookup(qty, info["tiers"])
            rows.append(
                {
                    "산식유형": f"구간별단가({label})",
                    "대상": store_names.get(store, store),
                    "현재수량": qty,
                    "현재단가": 현재단가,
                    "다음구간까지": (다음문턱값 - qty) if 다음문턱값 is not None else None,
                    "다음구간단가": 다음단가,
                }
            )

    linked_rules = [r for r in rules if _infer_규칙(r)["유형"] == "연동형"]
    if linked_rules and not agm_lv.empty:
        info = _infer_규칙(linked_rules[0])
        트리거제품군 = info.get("트리거제품군")
        qty_by_store = trigger_qty_by_store.get(트리거제품군, pd.Series(dtype=float))
        store_names = data.drop_duplicates("대리점코드").set_index("대리점코드")["대리점명"]
        for _, store_row in agm_lv.iterrows():
            store = store_row["대리점코드"]
            qty = qty_by_store.get(store, 0.0)
            tiers = [(store_row["Lv1"], info["tier_rates"][0]), (store_row["Lv2"], info["tier_rates"][1]), (store_row["Lv3"], info["tier_rates"][2])]
            현재단가, 다음문턱값, 다음단가 = _tier_lookup(qty, tiers)
            rows.append(
                {
                    "산식유형": f"연동형({트리거제품군} Lv)",
                    "대상": store_names.get(store, store),
                    "현재수량": qty,
                    "현재단가":현재단가,
                    "다음구간까지": (다음문턱값 - qty) if 다음문턱값 is not None else None,
                    "다음구간단가": 다음단가,
                }
            )

    if not rows:
        return pd.DataFrame(columns=["산식유형", "대상", "현재수량", "현재단가", "다음구간까지", "다음구간단가"])
    df = pd.DataFrame(rows)
    return df.sort_values("다음구간까지", na_position="last").reset_index(drop=True)


def _insight_품목별고정단가(label: str, applied_support: pd.DataFrame) -> list[str]:
    rows = applied_support[applied_support["지원유형"] == label]
    lines = [f"▶ {label} (품목별고정단가)"]
    lines.append("  장점: 판매수량과 무관하게 대당 단가가 고정돼 있어, 대리점 입장에서 예측하기 쉽고 정산이 단순합니다.")
    lines.append("  단점: 수량 문턱값이 없어, 이미 많이 판매하는 대리점을 더 밀어붙이는 추가 구매 유인 효과는 약합니다.")
    if not rows.empty:
        총액 = rows["지원금액"].sum()
        대리점수 = rows["대리점코드"].nunique()
        평균 = 총액 / 대리점수 if 대리점수 else 0
        lines.append(f"  이번 달: {대리점수}개 대리점에 총 {총액:,.0f}원 지원(대리점당 평균 {평균:,.0f}원).")
    return lines


def _insight_구간별단가(label: str, threshold_proximity: pd.DataFrame) -> list[str]:
    rows = threshold_proximity[threshold_proximity["산식유형"] == f"구간별단가({label})"]
    lines = [f"▶ {label} (구간별단가-절대수량)"]
    lines.append("  장점: 판매수량 구간을 넘길 때마다 단가가 올라가므로, 문턱값 근처 대리점에게 추가 구매를 유도하는 효과가 있습니다.")
    lines.append("  단점: 이미 최고 구간을 넘긴 대리점에는 더 이상 추가 유인이 없고, 최저 구간에도 못 미친 대리점에는 아직 아무 효과가 없습니다.")
    if not rows.empty:
        총_대상 = len(rows)
        # "근접"은 다음 구간 문턱값의 10% 이내로 남은 경우로 정의한다(제품군마다 규모가
        # 달라 절대 수량 기준보다 상대 기준이 더 일관된다).
        다음문턱값 = rows["현재수량"] + rows["다음구간까지"]
        근접_mask = rows["다음구간까지"].notna() & (rows["다음구간까지"] <= (다음문턱값 * 0.1).clip(lower=1))
        최고구간_mask = rows["다음구간까지"].isna()
        미달_mask = (~최고구간_mask) & (rows["현재단가"] == 0)
        lines.append(
            f"  이번 달(전체 {총_대상}곳): 다음 구간까지 근접(10% 이내)한 대리점 {int(근접_mask.sum())}곳, "
            f"이미 최고 구간 도달 {int(최고구간_mask.sum())}곳, 최저 구간에도 못 미친 대리점 {int(미달_mask.sum())}곳."
        )
    return lines


def _insight_연동형(label: str, applied_support: pd.DataFrame, threshold_proximity: pd.DataFrame) -> list[str]:
    rows = applied_support[applied_support["지원유형"].str.startswith(label)]
    lines = [f"▶ {label} (연동형)"]
    lines.append(
        "  장점: 트리거 제품군(예: AGM) 판매와 목표 품목 판매를 함께 늘리는 교차 판매 유인 효과가 있고, "
        "트리거 제품군 자체에도 지원이 붙는 경우 유인이 이중으로 강해집니다."
    )
    lines.append("  단점: 대형 대리점일수록 Lv 문턱값을 쉽게 넘겨 지원금이 소수 대형 대리점에 집중되는 경향이 있습니다.")
    if not rows.empty:
        by_store = rows.groupby("대리점코드")["지원금액"].sum().sort_values(ascending=False)
        총액 = by_store.sum()
        상위3 = by_store.head(3).sum()
        비중 = (상위3 / 총액 * 100) if 총액 else 0
        lines.append(f"  이번 달: 총 {len(by_store)}개 대리점 중 상위 3개 대리점이 전체 지원금액의 {비중:.0f}%를 차지합니다.")
    return lines


def _insight_무상증정(label: str, applied_support: pd.DataFrame) -> list[str]:
    rows = applied_support[applied_support["지원유형"] == label]
    lines = [f"▶ {label} (무상증정)"]
    lines.append("  장점: 일정 수량 이상 구매하면 자동으로 추가 지원이 붙어 대량 구매를 유도하는 효과가 뚜렷합니다.")
    lines.append("  단점: 문턱값 바로 아래(예: 1~2대 부족)에서 끊기는 대리점은 지원을 전혀 못 받아 형평성 불만이 생길 수 있습니다.")
    if not rows.empty:
        총액 = rows["지원금액"].sum()
        대리점수 = rows["대리점코드"].nunique()
        lines.append(f"  이번 달: {대리점수}개 대리점에 총 {총액:,.0f}원 지원.")
    return lines


def _build_promotion_insight(
    rules: list[dict], applied_support: pd.DataFrame, threshold_proximity: pd.DataFrame
) -> list[str]:
    """산식유형별로 이 프로모션 설계의 구조적 장단점과, 이번 달 실제 데이터가 보여주는
    신호를 함께 정리한다. 차기 프로모션 설계안을 시스템이 대신 결정하지는 않는다(PRD 3절
    비목표) — 팀장의 판단을 돕는 참고 의견일 뿐, 결론이 아니다."""
    label_유형: dict[str, str] = {}
    for rule in rules:
        label_유형.setdefault(rule["프로모션명"], rule["산식유형"])

    lines: list[str] = []
    for label, 유형 in label_유형.items():
        if 유형 == "품목별고정단가":
            lines.extend(_insight_품목별고정단가(label, applied_support))
        elif 유형 == "구간별단가-절대수량":
            lines.extend(_insight_구간별단가(label, threshold_proximity))
        elif 유형 == "연동형":
            lines.extend(_insight_연동형(label, applied_support, threshold_proximity))
        elif 유형 == "무상증정":
            lines.extend(_insight_무상증정(label, applied_support))
    return lines


def build_sales_like_table(workbook: str | Path | BinaryIO, 기준연월: str) -> IngestResult:
    """워크북을 읽어 대리점 단위 결과를 만든다. "DC율" 시트가 있어야 하는
    유일한 형식이다(`.docs/19_최종백데이터_전환_계획서.md`).

    `IngestResult.sales_like` 컬럼: 기준연월, 대리점코드, 대리점명, 판매수량, 매출액,
    전월매출, 매출증감율, 전년동월매출, 전년동월매출증감율, 영업이익, 영업이익율,
    매출DC율, 성장DC율, DC율, DC지원금액, 프로모션지원금액, 총지원금액, 원가율.
    전월매출·전년동월매출은 "DC율" 시트의
    "매출액" 병합 슈퍼헤더(하위 열이 "26년 8월"처럼 상대 연월 라벨이라 위치로 찾음)가
    있어야 채워진다 — 없으면 NaN이고 경고가 남는다.

    `IngestResult.applied_support`는 지원유형별(매출DC/성장DC/그 외 산식유형별 프로모션명)
    지원금액을 롱포맷으로 담은 DataFrame이다(0원인 건은 제외).

    `IngestResult.product_group_support`는 대리점 구분 없이 제품구분1(GB/AGM/PT 등)별
    판매수량·프로모션 지원금액 전사 합계다.

    `IngestResult.promotion_notice`는 프로모션명별로 "총 n개 품목" 수준으로 요약한
    안내 문구다.

    `IngestResult.threshold_proximity`는 구간별단가·연동형 규칙의 다음 구간까지
    남은 수량을 대상별로 보여준다.

    `IngestResult.promotion_insight`는 산식유형별 구조적 장단점 + 이번 달 데이터가
    보여주는 신호를 함께 정리한 참고 의견이다(차기 설계안을 대신 결정하지 않음,
    PRD 3절 비목표).
    """
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=False)
    if "DC율" not in wb.sheetnames:
        raise ColumnNotFoundError('워크북에 "DC율" 시트가 없습니다 — 지원하는 형식이 아닙니다.')

    warnings: list[str] = []

    dc율 = extract_DC율(wb)
    제품마스터, rules, agm_lv = _extract_프로모션기준(wb["프로모션 기준"])
    # 프로모션 기준(제품마스터)에 없는 제품코드(예: 이륜용 배터리)는 이 집계 대상이
    # 아니라서 조용히 제외한다 — 매번 뜨는 안내가 오히려 노이즈라는 사용자 피드백으로
    # 화면 경고는 없앴다(제외된 코드 자체는 `extract_DATA`가 여전히 돌려주므로 필요하면
    # 디버깅에 쓸 수 있다).
    data, excluded_codes = extract_DATA(wb, set(제품마스터["제품코드"]))

    data = data.merge(제품마스터[["제품코드", "제품구분1", "기준가", "DC미적용"]], on="제품코드", how="left")
    data = data.merge(dc율[["대리점코드", "매출DC율", "성장DC율"]], on="대리점코드", how="left")
    data[["매출DC율", "성장DC율"]] = data[["매출DC율", "성장DC율"]].fillna(0.0)

    # --- DC지원금액 = 기준가 × 판매수량 × DC율 (PRD 6.3.3), "DC율 미적용" 제품은 0원 ---
    eligible = ~data["DC미적용"].fillna(False)
    data["매출DC금액"] = 0.0
    data["성장DC금액"] = 0.0
    data.loc[eligible, "매출DC금액"] = data.loc[eligible, "기준가"] * data.loc[eligible, "판매수량"] * data.loc[eligible, "매출DC율"]
    data.loc[eligible, "성장DC금액"] = data.loc[eligible, "기준가"] * data.loc[eligible, "판매수량"] * data.loc[eligible, "성장DC율"]

    # --- 프로모션 지원금액: 산식유형별로 계산 ---
    data["프로모션금액"] = 0.0
    support_by_label: dict[str, pd.Series] = {}
    trigger_qty_by_store: dict[str, pd.Series] = {}
    # 연동형(예: Level up)은 대상 품목(GB 등)뿐 아니라 트리거제품군(AGM) 자체도 Lv 달성별로
    # 대당 지원금액을 받는다("트리거지원금액" 컬럼). 같은 트리거제품군·단가표를 쓰는 연동형
    # 규칙이 대상 품목 수만큼(예: GB 17개 품목) 반복되므로, 트리거 자체지원은 유니크한
    # (트리거제품군, 단가표) 조합별로 한 번만 적용해야 중복 합산되지 않는다.
    linked_trigger_configs: dict[tuple[str, tuple[float, ...]], str] = {}

    def _qty_by_group(group: str) -> pd.Series:
        if group not in trigger_qty_by_store:
            trigger_qty_by_store[group] = data.loc[data["제품구분1"] == group].groupby("대리점코드")["판매수량"].sum()
        return trigger_qty_by_store[group]

    def _rate_by_store_for_tiers(qty_by_store: pd.Series, lv_idx: pd.DataFrame, tier_rates: list[float]) -> pd.Series:
        def _rate(store: str) -> float:
            if store not in lv_idx.index:
                return 0.0
            qty = qty_by_store.get(store, 0.0)
            tiers = [
                (lv_idx.loc[store, "Lv1"], tier_rates[0]),
                (lv_idx.loc[store, "Lv2"], tier_rates[1]),
                (lv_idx.loc[store, "Lv3"], tier_rates[2]),
            ]
            return _tier_lookup(qty, tiers)[0]

        return pd.Series({store: _rate(store) for store in lv_idx.index})

    # 구간별단가(절대수량, 예: "특화")는 프로모션명 그룹에 속한 "모든" 제품코드의
    # 판매수량을 대리점 단위로 합산한 값으로 구간을 판정한 뒤, 그 구간 단가를 그룹
    # 안의 각 제품 판매수량에 곱한다(개별 제품코드 단위로 판정하면 틀림 —
    # `.docs/19_최종백데이터_전환_계획서.md` 실사례 검증: 대형 대리점 6곳을 공식
    # "26년 8월 DC 총괄표.xlsx"와 대조해 이 규칙을 확인했다. AGM Level up의 트리거
    # 판정과 같은 원리다). 아래 메인 루프에서는 처리하지 않고 별도 그룹 단위로
    # 처리한다.
    tiered_group_codes: dict[str, list[dict]] = {}
    for rule in rules:
        if _infer_규칙(rule)["유형"] == "구간별단가-절대수량":
            tiered_group_codes.setdefault(rule["프로모션명"], []).append(rule)

    for rule in rules:
        info = _infer_규칙(rule)
        label = rule["프로모션명"]
        if info["유형"] == "품목별고정단가":
            mask = data["제품코드"] == rule["제품코드"]
            if not mask.any():
                continue
            amount = data.loc[mask, "판매수량"] * info["단가"]
        elif info["유형"] == "구간별단가-절대수량":
            continue  # 아래 별도 그룹 단위 패스에서 한 번에 처리한다.
        elif info["유형"] == "무상증정":
            mask = data["제품코드"] == rule["제품코드"]
            if not mask.any():
                continue
            amount = (data.loc[mask, "판매수량"] // info["n"]) * data.loc[mask, "기준가"]
        elif info["유형"] == "연동형":
            트리거 = info.get("트리거제품군")
            mask = data["제품코드"] == rule["제품코드"]
            if not 트리거 or agm_lv.empty:
                warnings.append(f'"{label}"(연동형) 계산에 필요한 트리거제품군 또는 대리점별 문턱값 표를 찾지 못했습니다.')
                continue
            if not mask.any():
                continue
            qty_by_store = _qty_by_group(트리거)
            lv_idx = agm_lv.set_index("대리점코드")
            rate_by_store = _rate_by_store_for_tiers(qty_by_store, lv_idx, info["tier_rates"])
            amount = data.loc[mask, "판매수량"] * data.loc[mask, "대리점코드"].map(rate_by_store).fillna(0.0)

            트리거단가 = tuple(info.get("트리거_tier_rates") or (0, 0, 0))
            if any(트리거단가):
                linked_trigger_configs.setdefault((트리거, 트리거단가), label)
        else:
            warnings.append(f'"{label}"의 산식유형("{rule["산식유형"]}")을 알아보지 못해 계산하지 않았습니다.')
            continue

        data.loc[mask, "프로모션금액"] += amount
        by_store = data.loc[mask].assign(_금액=amount).groupby("대리점코드")["_금액"].sum()
        support_by_label[label] = support_by_label.get(label, pd.Series(dtype=float)).add(by_store, fill_value=0.0)

    # --- 구간별단가(절대수량) 그룹 단위 계산: 프로모션명 그룹의 합산 판매수량으로
    # 구간을 판정하고, 그 구간 단가를 그룹 안의 각 제품 판매수량에 곱한다 ---
    for label, group_rules in tiered_group_codes.items():
        info = _infer_규칙(group_rules[0])
        if not info["tiers"]:
            continue
        codes = [r["제품코드"] for r in group_rules]
        mask = data["제품코드"].isin(codes)
        if not mask.any():
            continue
        qty_by_store = data.loc[mask].groupby("대리점코드")["판매수량"].sum()
        rate_by_store = qty_by_store.apply(lambda q: _tier_lookup(q, info["tiers"])[0])
        amount = data.loc[mask, "판매수량"] * data.loc[mask, "대리점코드"].map(rate_by_store).fillna(0.0)
        data.loc[mask, "프로모션금액"] += amount
        by_store = data.loc[mask].assign(_금액=amount).groupby("대리점코드")["_금액"].sum()
        support_by_label[label] = support_by_label.get(label, pd.Series(dtype=float)).add(by_store, fill_value=0.0)

    # --- 연동형 트리거제품군(예: AGM) 자체 지원: 대상 품목 규칙 수만큼이 아니라
    # 유니크한 (트리거제품군, 단가표) 조합별로 딱 한 번만 적용한다 ---
    for (트리거, 트리거단가), 대표라벨 in linked_trigger_configs.items():
        if agm_lv.empty:
            continue
        qty_by_store = _qty_by_group(트리거)
        lv_idx = agm_lv.set_index("대리점코드")
        rate_by_store = _rate_by_store_for_tiers(qty_by_store, lv_idx, list(트리거단가))
        mask = data["제품구분1"] == 트리거
        if not mask.any():
            continue
        amount = data.loc[mask, "판매수량"] * data.loc[mask, "대리점코드"].map(rate_by_store).fillna(0.0)
        data.loc[mask, "프로모션금액"] += amount
        label = f"{대표라벨}(트리거 {트리거} 자체지원)"
        by_store = data.loc[mask].assign(_금액=amount).groupby("대리점코드")["_금액"].sum()
        support_by_label[label] = support_by_label.get(label, pd.Series(dtype=float)).add(by_store, fill_value=0.0)

    store_agg = data.groupby("대리점코드", as_index=False).agg(
        대리점명=("대리점명", "first"),
        매출액=("매출액", "sum"),
        판매수량=("판매수량", "sum"),
        영업이익=("영업이익", "sum"),
        매출원가=("매출원가", "sum"),
        매출DC금액=("매출DC금액", "sum"),
        성장DC금액=("성장DC금액", "sum"),
        프로모션금액=("프로모션금액", "sum"),
    )
    store_agg = store_agg.merge(
        dc율[["대리점코드", "매출DC율", "성장DC율", "전월매출", "전년동월매출"]], on="대리점코드", how="left"
    )

    out = store_agg.copy()
    out["기준연월"] = 기준연월
    out["DC지원금액"] = out["매출DC금액"] + out["성장DC금액"]
    out["프로모션지원금액"] = out["프로모션금액"]
    out["총지원금액"] = out["DC지원금액"] + out["프로모션지원금액"]
    out["영업이익율"] = (out["영업이익"] / out["매출액"].replace(0, pd.NA)) * 100
    out["원가율"] = (out["매출원가"] / out["매출액"].replace(0, pd.NA)) * 100
    out["매출증감율"] = ((out["매출액"] - out["전월매출"]) / out["전월매출"].replace(0, pd.NA)) * 100
    out["전년동월매출증감율"] = ((out["매출액"] - out["전년동월매출"]) / out["전년동월매출"].replace(0, pd.NA)) * 100
    # "DC율"은 미적용 제품을 뺀 실제 평균 적용률이라, DC율 시트의 대리점별 계약
    # 요율(매출DC율/성장DC율, 아래서 %로 변환)과는 다를 수 있다 — 의도된 차이다.
    out["DC율"] = (out["DC지원금액"] / out["매출액"].replace(0, pd.NA)) * 100
    out["매출DC율"] = out["매출DC율"] * 100
    out["성장DC율"] = out["성장DC율"] * 100

    if out["전월매출"].isna().all():
        warnings.append("전월매출이 이 워크북에 없어 매출증감율을 계산할 수 없습니다(DC율 시트의 \"매출액\" 병합 헤더를 확인하세요).")
    if out["전년동월매출"].isna().all():
        warnings.append("전년동월매출이 이 워크북에 없어 전년동월매출증감율을 계산할 수 없습니다(DC율 시트의 \"매출액\" 병합 헤더를 확인하세요).")

    ordered = out[
        [
            "기준연월", "대리점코드", "대리점명", "판매수량", "매출액", "전월매출", "매출증감율",
            "전년동월매출", "전년동월매출증감율",
            "영업이익", "영업이익율", "매출DC율", "성장DC율", "DC율", "DC지원금액",
            "프로모션지원금액", "총지원금액", "원가율",
        ]
    ]

    applied_support_parts = [
        _support_rows(store_agg, 기준연월, "매출DC", store_agg["매출DC금액"]),
        _support_rows(store_agg, 기준연월, "성장DC", store_agg["성장DC금액"]),
    ]
    for label, amounts in support_by_label.items():
        store_view = store_agg.set_index("대리점코드")
        aligned = store_agg["대리점코드"].map(amounts).fillna(0.0)
        applied_support_parts.append(_support_rows(store_agg, 기준연월, label, aligned))
    applied_support = pd.concat(applied_support_parts, ignore_index=True)

    product_group_support = (
        data.groupby("제품구분1", as_index=False)
        .agg(판매수량=("판매수량", "sum"), 지원금액=("프로모션금액", "sum"))
        .rename(columns={"제품구분1": "제품군"})
    )
    product_group_support.insert(0, "기준연월", 기준연월)

    promotion_notice = _build_promotion_notice(rules)
    threshold_proximity = _build_threshold_proximity(data, rules, agm_lv, trigger_qty_by_store)
    promotion_insight = _build_promotion_insight(rules, applied_support, threshold_proximity)

    return IngestResult(
        sales_like=ordered,
        applied_support=applied_support,
        product_group_support=product_group_support,
        promotion_notice=promotion_notice,
        threshold_proximity=threshold_proximity,
        promotion_insight=promotion_insight,
        warnings=warnings,
    )
