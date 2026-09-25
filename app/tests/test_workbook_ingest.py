from __future__ import annotations

import io

import openpyxl
import pandas as pd
import pytest

from app.workbook_ingest import ColumnNotFoundError, build_sales_like_table

# ---------------------------------------------------------------------------
# DC율/프로모션 기준/DATA 시트로 구성된 실데이터 워크북 경로
# (`.docs/17_실데이터_워크북_반영_계획서.md`·`19_최종백데이터_전환_계획서.md`) — 손으로
# 검산 가능한 합성 워크북으로 세 산식(품목별고정단가/구간별단가-절대수량/연동형)과
# DC 미적용 처리, 프로모션 기준에 없는 제품코드 제외를 각각 검증한다.
#
# 산식은 "산식유형"·"트리거제품군" 컬럼(고정 어휘)만으로 분기해야 하고, 프로모션명은
# 화면 표시용 자유 텍스트일 뿐이어야 한다 — 그래서 프로모션명은 일부러 산식유형과
# 무관한 이름("대형 특가", "특화 캠페인", "여름 Lv 이벤트")을 붙여 이름이 계산에
# 영향을 주지 않는지 검증한다.
# ---------------------------------------------------------------------------


def _write_workbook_from_grids(sheets: dict[str, list[list]]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, grid in sheets.items():
        ws = wb.create_sheet(name)
        for row in grid:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _write_dc율_워크북() -> io.BytesIO:
    dc율 = [
        ["■ 대리점별 DC율"],
        ["거래처코드", "거래처명", "매출DC", "성장DC"],
        ["S1", "테스트대리점1", 0.05, 0.02],
        ["S2", "테스트대리점2", 0.10, 0.00],
    ]
    # 헤더 컬럼(0-based): 0 구분, 1 제품코드, 2 제품명, 3 제품구분1, 4 제품구분2, 5 기준가,
    # 6 비고, 7 프로모션명, 8 시작일, 9 종료일, 10 지원내용, 11-13 지원 조건, 14-16 지원 금액,
    # 17 (미사용), 18 코드, 19 대리점, 20 AGM Lv.1, 21 AGM Lv.2, 22 AGM Lv.3,
    # 23 산식유형, 24 트리거제품군, 25-27 트리거제품군 지원금액(Lv.1/2/3)
    프로모션기준 = [
        [
            "구분", "제품코드", "제품명", "제품구분1", "제품구분2", "기준가", "비고",
            "프로모션명", "시작일", "종료일", "지원내용",
            "지원 조건", None, None, "지원 금액", None, None,
            None, "코드", "대리점", "AGM Lv.1", "AGM Lv.2", "AGM Lv.3",
            "산식유형", "트리거제품군", "트리거제품군 지원금액", None, None,
        ],
        [
            1, "CODE_DC", "Product DC", "GB", "GB 중형", 1000, None,
            None, None, None, None, None, None, None, None, None, None,
            None, "S1", "테스트대리점1", 100, 200, 300,
        ],
        [
            2, "CODE_NODC", "Product NoDC", "GB", "GB 중형", 2000, "DC율 미적용",
            None, None, None, None, None, None, None, None, None, None,
            None, "S2", "테스트대리점2", 50, 100, 150,
        ],
        [
            3, "CODE_LARGE", "Product Large", "GB", "GB 대형", 500, None,
            "대형 특가", "2026-08-01", "2026-08-31", "대당 금액 지원",
            None, None, None, 3000, None, None,
            None, None, None, None, None, None,
            "품목별고정단가", None,
        ],
        # CODE_SPECIAL: "특화"(구간별단가)와 "대형"(품목별고정단가, 다른 단가 4000)이
        # 같은 제품코드에 동시 적용되는 실제 사례(GB73010류)를 재현 — 둘 다 합산돼야 한다.
        [
            4, "CODE_SPECIAL", "Product Special", "GB", "GB 대형", 800, None,
            "특화 캠페인", "2026-08-01", "2026-08-31", "대당 금액 지원",
            "50대 이상", "100대 이상", "300대 이상", 2000, 3000, 5000,
            None, None, None, None, None, None,
            "구간별단가-절대수량", None,
        ],
        [
            41, "CODE_SPECIAL", "Product Special", "GB", "GB 대형", 800, None,
            "대형 특가", "2026-08-01", "2026-08-31", "대당 금액 지원",
            None, None, None, 4000, None, None,
            None, None, None, None, None, None,
            "품목별고정단가", None,
        ],
        [5, "CODE_AGM", "AGM Product", "AGM", "AGM", 100, None],
        [
            6, "CODE_LEVELUP", "Level Up GB", "GB", "GB 중형", 1500, None,
            "여름 Lv 이벤트", "2026-08-01", "2026-08-31", "AGM Lv 달성별 대당 금액 지원",
            "AGM Lv.1", "AGM Lv.2", "AGM Lv.3", 1000, 2000, 3000,
            None, None, None, None, None, None,
            "연동형", "AGM", 3000, 4000, 5000,
        ],
        # CODE_LEVELUP2: 같은 트리거(AGM)·같은 트리거지원금액표를 쓰는 두 번째 대상
        # 품목 — 트리거(AGM) 자체지원이 대상 품목 수만큼(2번) 중복 합산되면 안 된다.
        [
            7, "CODE_LEVELUP2", "Level Up GB 2", "GB", "GB 중형", 1200, None,
            "여름 Lv 이벤트", "2026-08-01", "2026-08-31", "AGM Lv 달성별 대당 금액 지원",
            "AGM Lv.1", "AGM Lv.2", "AGM Lv.3", 1000, 2000, 3000,
            None, None, None, None, None, None,
            "연동형", "AGM", 3000, 4000, 5000,
        ],
    ]
    data = [
        [
            "거래처코드", "대리점명", "제품코드", "제품명", "구분", "판매구분",
            "제품구분1", "제품구분2", "제품구분3", "매출액(Total)", "매출수량(SET)",
            "영업이익(A)", "매출원가(S)Tot", "용량(단위당)",
        ],
        ["S1", "테스트대리점1", "CODE_DC", "Product DC", "대리점", "차량용", "GB", "GB 중형", "GB 중형", 100000, 10, 10000, 70000, 40],
        ["S1", "테스트대리점1", "CODE_NODC", "Product NoDC", "대리점", "차량용", "GB", "GB 중형", "GB 중형", 50000, 5, 5000, 35000, 40],
        ["S1", "테스트대리점1", "CODE_LARGE", "Product Large", "대리점", "차량용", "GB", "GB 대형", "GB 대형", 200000, 20, 20000, 140000, 120],
        ["S1", "테스트대리점1", "CODE_SPECIAL", "Product Special", "대리점", "차량용", "GB", "GB 대형", "GB 대형", 960000, 120, 96000, 672000, 120],
        ["S1", "테스트대리점1", "CODE_AGM", "AGM Product", "대리점", "차량용", "AGM", "AGM", "AGM", 250000, 250, 25000, 175000, 70],
        ["S1", "테스트대리점1", "CODE_LEVELUP", "Level Up GB", "대리점", "차량용", "GB", "GB 중형", "GB 중형", 120000, 8, 12000, 84000, 70],
        ["S1", "테스트대리점1", "CODE_LEVELUP2", "Level Up GB 2", "대리점", "차량용", "GB", "GB 중형", "GB 중형", 60000, 4, 6000, 42000, 70],
        ["S1", "테스트대리점1", "QMS_TEST", "이륜용 배터리", "대리점", "이륜용", "이륜용", "이륜용", "이륜용", 9999, 3, 999, 9000, 12],
        ["S2", "테스트대리점2", "CODE_AGM", "AGM Product", "대리점", "차량용", "AGM", "AGM", "AGM", 120000, 120, 12000, 84000, 70],
        ["S2", "테스트대리점2", "CODE_LEVELUP", "Level Up GB", "대리점", "차량용", "GB", "GB 중형", "GB 중형", 40000, 4, 4000, 28000, 70],
    ]
    return _write_workbook_from_grids({"DC율": dc율, "프로모션 기준": 프로모션기준, "DATA": data})


class TestDC율기반_실데이터워크북:
    @pytest.fixture()
    def result(self):
        buf = _write_dc율_워크북()
        return build_sales_like_table(buf, "2026-08")

    def test_columns(self, result):
        assert list(result.sales_like.columns) == [
            "기준연월", "대리점코드", "대리점명", "판매수량", "매출액", "전월매출",
            "매출증감율", "전년동월매출", "전년동월매출증감율",
            "영업이익", "영업이익율", "매출DC율", "성장DC율", "DC율", "DC지원금액",
            "프로모션지원금액", "총지원금액", "원가율",
        ]

    def test_excludes_unknown_product_code_silently(self, result):
        # 프로모션 기준(제품마스터)에 없는 제품코드(예: 이륜용)는 조용히 제외한다 —
        # 매번 뜨는 안내가 노이즈라는 사용자 피드백으로 화면 경고는 없앴다.
        assert "QMS_TEST" not in result.sales_like.to_string()
        assert not any("QMS_TEST" in w for w in result.warnings)

    def test_영업이익_원가_grabbed_directly_no_gap_warning(self, result):
        # 이 형식은 DATA에 영업이익·매출원가가 실제로 있어 "없어서 계산 못 함"
        # 경고가 나오면 안 된다.
        assert not any("영업이익" in w and "계산할 수 없" in w for w in result.warnings)
        s1 = result.sales_like.set_index("대리점코드").loc["S1"]
        assert s1["영업이익율"] == pytest.approx(10.0)
        assert s1["원가율"] == pytest.approx(70.0)

    def test_dc_amount_computed_per_line_excluding_미적용(self, result):
        # S1: (CODE_DC 10000 + CODE_LARGE 10000 + CODE_SPECIAL 96000 + CODE_AGM 25000
        #      + CODE_LEVELUP 12000 + CODE_LEVELUP2 4800) × (0.05+0.02) = 157800 × 0.07 = 11046
        #      (CODE_NODC는 DC율 미적용이라 기준가×수량 10000이 제외된다)
        s1 = result.sales_like.set_index("대리점코드").loc["S1"]
        assert s1["DC지원금액"] == pytest.approx(11046.0)
        support = result.applied_support.set_index(["대리점코드", "지원유형"])["지원금액"]
        assert support[("S1", "매출DC")] == pytest.approx(7890.0)
        assert support[("S1", "성장DC")] == pytest.approx(3156.0)

    def test_flat_and_tiered_promotion_driven_by_산식유형_not_프로모션명(self, result):
        # 프로모션명은 "대형 특가"/"특화 캠페인"처럼 산식유형과 무관한 이름이지만,
        # 산식유형 컬럼(품목별고정단가/구간별단가-절대수량)만으로 정확히 계산돼야 한다.
        # CODE_SPECIAL은 "특화"(구간별단가)와 "대형"(품목별고정단가, 4,000원)이 같은
        # 제품코드에 동시 적용되는 실제 사례(GB73010류)라 "대형 특가" 라벨 합계에
        # CODE_LARGE(20대×3,000원)와 CODE_SPECIAL(120대×4,000원)이 함께 잡혀야 한다.
        support = result.applied_support.set_index(["대리점코드", "지원유형"])["지원금액"]
        assert support[("S1", "대형 특가")] == pytest.approx(60000.0 + 480000.0)
        assert support[("S1", "특화 캠페인")] == pytest.approx(360000.0)  # 120대(100대 이상 구간) × 3,000원

    def test_linked_agm_uses_store_specific_threshold(self, result):
        # S1: AGM 총 판매수량 250 → Lv1=100,Lv2=200,Lv3=300 중 [200,300) 구간 = Lv.2 → 2,000원
        #     대상 품목 2개(CODE_LEVELUP 8대, CODE_LEVELUP2 4대) 모두 같은 "여름 Lv 이벤트"
        #     라벨이라 (8+4)×2,000원 = 24,000원으로 합산된다.
        # S2: AGM 총 판매수량 120 → Lv1=50,Lv2=100,Lv3=150 중 [100,150) 구간 = Lv.2 → 2,000원
        #     CODE_LEVELUP 4대 × 2,000원 = 8,000원 (CODE_LEVELUP2는 S2에 판매 없음)
        support = result.applied_support.set_index(["대리점코드", "지원유형"])["지원금액"]
        assert support[("S1", "여름 Lv 이벤트")] == pytest.approx(24000.0)
        assert support[("S2", "여름 Lv 이벤트")] == pytest.approx(8000.0)

    def test_linked_trigger_group_self_support_applied_once_not_per_target_product(self, result):
        # 트리거제품군(AGM) 자체도 Lv 달성별 대당 지원금액(3,000~5,000원)을 받는다.
        # CODE_LEVELUP·CODE_LEVELUP2 둘 다 같은 트리거(AGM)·같은 단가표를 쓰므로,
        # 대상 품목이 2개라고 해서 AGM 자체지원이 2번 중복 합산되면 안 된다.
        # S1: AGM 250대, Lv.2 구간(200~300) → 트리거단가 4,000원 → 250×4,000=1,000,000원
        # S2: AGM 120대, Lv.2 구간(100~150, S2 고유 문턱값) → 4,000원 → 120×4,000=480,000원
        support = result.applied_support.set_index(["대리점코드", "지원유형"])["지원금액"]
        assert support[("S1", "여름 Lv 이벤트(트리거 AGM 자체지원)")] == pytest.approx(1000000.0)
        assert support[("S2", "여름 Lv 이벤트(트리거 AGM 자체지원)")] == pytest.approx(480000.0)

    def test_total_support_equals_dc_plus_promotion(self, result):
        df = result.sales_like
        assert (df["총지원금액"] - (df["DC지원금액"] + df["프로모션지원금액"])).abs().max() < 1e-6

    def test_promotion_insight_covers_every_프로모션명(self, result):
        joined = " ".join(result.promotion_insight)
        assert "▶ 대형 특가 (품목별고정단가)" in result.promotion_insight
        assert "▶ 특화 캠페인 (구간별단가-절대수량)" in result.promotion_insight
        assert "▶ 여름 Lv 이벤트 (연동형)" in result.promotion_insight
        assert "장점" in joined and "단점" in joined

    def test_promotion_insight_fixed_rate_uses_total_지원금액_and_대리점수(self, result):
        # "대형 특가"는 S1만 받는다(60,000+480,000=540,000원, 1개 대리점).
        idx = result.promotion_insight.index("▶ 대형 특가 (품목별고정단가)")
        summary_line = result.promotion_insight[idx + 3]
        assert "1개 대리점" in summary_line
        assert "540,000원" in summary_line

    def test_promotion_insight_linked_concentration(self, result):
        # 연동형 지원(대상+트리거 자체지원)을 받는 대리점은 S1·S2 2곳뿐이라, 상위 3개
        # 대리점 비중은 사실상 전체(100%)와 같다.
        idx = result.promotion_insight.index("▶ 여름 Lv 이벤트 (연동형)")
        summary_line = result.promotion_insight[idx + 3]
        assert "2개 대리점" in summary_line
        assert "100%" in summary_line

    def test_no_전월매출_warns_but_does_not_fail(self, result):
        assert result.sales_like["전월매출"].isna().all()
        assert result.sales_like["매출증감율"].isna().all()
        assert any("전월매출" in w for w in result.warnings)

    def test_product_group_support_single_제품구분1_axis(self, result):
        # 제품군별 요약은 이제 제품구분1(GB/AGM) 단일 축이다 — 프로모션 종류별
        # 내역은 이미 지원유형별 지원금액 내역(applied_support)에 있으므로 중복이 아니다.
        groups = set(result.product_group_support["제품군"])
        assert groups == {"GB", "AGM"}

    def test_threshold_proximity_covers_구간별단가와_연동형(self, result):
        proximity = result.threshold_proximity
        assert set(proximity["산식유형"]) == {"구간별단가(특화 캠페인)", "연동형(AGM Lv)"}
        # S1은 AGM 250대로 Lv.2(200) 도달, Lv.3(300)까지 50대 남음
        s1_agm = proximity[(proximity["산식유형"] == "연동형(AGM Lv)") & (proximity["대상"] == "테스트대리점1")].iloc[0]
        assert s1_agm["다음구간까지"] == pytest.approx(50.0)
        assert s1_agm["다음구간단가"] == pytest.approx(3000.0)

    def test_promotion_notice_shows_all_flat_rate_variants(self, result):
        # "대형 특가" 프로모션명 그룹 안에 3,000원(CODE_LARGE)과 4,000원(CODE_SPECIAL)
        # 두 단가가 섞여 있다 — 대표 한 줄(3,000원)만 보여주면 안 되고 둘 다 보여야 한다.
        joined = " ".join(result.promotion_notice)
        assert "3,000원" in joined and "4,000원" in joined
        assert "1개 품목" in joined  # 각각 1개 품목이라는 표시

    def test_promotion_notice_mentions_trigger_self_support(self, result):
        joined = " ".join(result.promotion_notice)
        assert "AGM 자체도 Lv 달성 시 대당 지원" in joined
        assert "3,000원" in joined and "5,000원" in joined


def _write_dc율_워크북_with_매출액이력(rows: list[dict]) -> io.BytesIO:
    """실제 파일처럼 "DC율" 시트에 "매출액" 병합 슈퍼헤더(2행 헤더, 하위 라벨은
    "26년 8월"처럼 상대 연월 문자열)가 있는 워크북을 만든다. 최소한의 "프로모션 기준"
    (산식유형 없음, 규칙 없음)·"DATA" 시트만 곁들여 build_sales_like_table 전체가
    돈다는 것까지 확인한다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("DC율")
    ws.append(["■ 대리점별 DC율"])
    ws.append(["거래처코드", "거래처명", "매출DC", "성장DC", "매출액", None, None, None, None])
    ws.merge_cells("A2:A3")
    ws.merge_cells("B2:B3")
    ws.merge_cells("C2:C3")
    ws.merge_cells("D2:D3")
    ws.merge_cells("E2:I2")
    ws.append([None, None, None, None, "26년 8월", "26년 7월", "전월 대비 증감율", "25년 8월", "전년 동월 대비 증감율"])
    for row in rows:
        ws.append(
            [
                row["대리점코드"], row["대리점명"], row["매출DC"], row["성장DC"],
                row["이번달매출"], row["전월매출"], row["전월대비증감율"], row["전년동월매출"], row["전년동월대비증감율"],
            ]
        )

    promo_ws = wb.create_sheet("프로모션 기준")
    promo_ws.append(["제품코드", "제품명", "제품구분1", "기준가", "프로모션명", "지원 조건", None, None, "지원 금액", "산식유형"])
    for row in rows:
        promo_ws.append([f"CODE_{row['대리점코드']}", "Product", "GB", 1000])

    data_ws = wb.create_sheet("DATA")
    data_ws.append(["거래처코드", "대리점명", "제품코드", "매출액(Total)", "매출수량(SET)", "영업이익(A)", "매출원가(S)Tot"])
    for row in rows:
        data_ws.append([row["대리점코드"], row["대리점명"], f"CODE_{row['대리점코드']}", row["이번달매출"], 1, 0, 0])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestDC율_매출액_병합헤더:
    """실제 업로드 파일처럼 "매출액" 병합 슈퍼헤더 아래 "26년 8월"·"26년 7월"·
    "25년 8월"(상대 연월, 고정 텍스트 아님) 하위 라벨로 전월매출·전년동월매출이
    들어온 경우를 검증한다."""

    @pytest.fixture()
    def result(self):
        buf = _write_dc율_워크북_with_매출액이력(
            [
                {
                    "대리점코드": "S1", "대리점명": "테스트대리점1", "매출DC": 0.05, "성장DC": 0.02,
                    "이번달매출": 100000, "전월매출": 80000, "전월대비증감율": 0.25,
                    "전년동월매출": 50000, "전년동월대비증감율": 1.0,
                }
            ]
        )
        return build_sales_like_table(buf, "2026-08")

    def test_전월매출_전년동월매출_추출됨(self, result):
        s1 = result.sales_like.set_index("대리점코드").loc["S1"]
        assert s1["전월매출"] == pytest.approx(80000.0)
        assert s1["전년동월매출"] == pytest.approx(50000.0)

    def test_매출증감율_계산됨(self, result):
        # 매출액(DATA 기준 합계, 100000) 대비 전월매출(80000): (100000-80000)/80000*100 = 25.0
        s1 = result.sales_like.set_index("대리점코드").loc["S1"]
        assert s1["매출증감율"] == pytest.approx(25.0)
        assert s1["전년동월매출증감율"] == pytest.approx((100000 - 50000) / 50000 * 100)

    def test_no_missing_data_warning(self, result):
        assert not any("전월매출이 이 워크북에 없어" in w for w in result.warnings)
        assert not any("전년동월매출이 이 워크북에 없어" in w for w in result.warnings)


def _write_구간별단가_그룹합산_워크북() -> io.BytesIO:
    """구간별단가-절대수량(예: "특화")은 같은 프로모션명 아래 제품코드가 여러 개면
    그 대리점의 **합산 판매수량**으로 구간을 판정해야 한다 — 실제 "26년 8월 DC
    총괄표.xlsx"와 대조해 확인된 사례다(`.docs/19_최종백데이터_전환_계획서.md`).
    CODE_A(40대, 개별로는 50대 미만이라 구간 미달)와 CODE_B(70대, 개별로는
    50~99대 구간)를 합치면 110대로 100~299대 구간에 들어간다 — 제품코드별로 각각
    판정하면 0원+140,000원=140,000원이 되지만, 합산 판정이 맞다면 (40+70)대 ×
    3,000원 = 330,000원이어야 한다."""
    dc율 = [
        ["거래처코드", "거래처명", "매출DC", "성장DC"],
        ["S1", "테스트대리점1", 0.0, 0.0],
    ]
    프로모션기준 = [
        [
            "제품코드", "제품명", "제품구분1", "기준가", "프로모션명",
            "지원 조건", None, None, "지원 금액", None, None, "산식유형",
        ],
        [
            "CODE_A", "Product A", "GB", 1000, "특화 캠페인",
            "50대 이상", "100대 이상", "300대 이상", 2000, 3000, 5000, "구간별단가-절대수량",
        ],
        [
            "CODE_B", "Product B", "GB", 1000, "특화 캠페인",
            "50대 이상", "100대 이상", "300대 이상", 2000, 3000, 5000, "구간별단가-절대수량",
        ],
    ]
    data = [
        ["거래처코드", "대리점명", "제품코드", "매출액(Total)", "매출수량(SET)", "영업이익(A)", "매출원가(A)Tot"],
        ["S1", "테스트대리점1", "CODE_A", 40000, 40, 4000, 28000],
        ["S1", "테스트대리점1", "CODE_B", 70000, 70, 7000, 49000],
    ]
    return _write_workbook_from_grids({"DC율": dc율, "프로모션 기준": 프로모션기준, "DATA": data})


class Test구간별단가_그룹합산판정:
    @pytest.fixture()
    def result(self):
        buf = _write_구간별단가_그룹합산_워크북()
        return build_sales_like_table(buf, "2026-08")

    def test_합산_판매수량으로_구간_판정(self, result):
        support = result.applied_support.set_index(["대리점코드", "지원유형"])["지원금액"]
        # 개별 판정(오답): CODE_A 40대→0원 + CODE_B 70대(50~99구간)→70×2,000=140,000원 = 140,000원
        # 합산 판정(정답): (40+70)=110대 → 100~299 구간 → 110×3,000 = 330,000원
        assert support[("S1", "특화 캠페인")] == pytest.approx(330000.0)

    def test_threshold_proximity_uses_합산_수량(self, result):
        proximity = result.threshold_proximity
        row = proximity[proximity["산식유형"] == "구간별단가(특화 캠페인)"].iloc[0]
        assert row["현재수량"] == pytest.approx(110.0)
        assert row["현재단가"] == pytest.approx(3000.0)
        assert row["다음구간까지"] == pytest.approx(190.0)  # 300 - 110
        assert row["다음구간단가"] == pytest.approx(5000.0)


class TestMissingFormulaTypeColumn:
    def test_raises_column_not_found_when_산식유형_missing(self):
        # "산식유형" 컬럼이 아예 없으면 옛 프로모션명 텍스트 추측으로 되돌아가지
        # 않고, 명확한 오류를 낸다(`.docs/19_최종백데이터_전환_계획서.md`).
        dc율 = [["거래처코드", "거래처명", "매출DC", "성장DC"], ["S1", "테스트대리점1", 0.05, 0.02]]
        프로모션기준 = [
            ["제품코드", "제품명", "제품구분1", "기준가", "프로모션명", "지원 조건", None, None, "지원 금액"],
            ["CODE_A", "Product A", "GB", 1000, "대형", None, None, None, 3000],
        ]
        data = [
            ["거래처코드", "대리점명", "제품코드", "매출액(Total)", "매출수량(SET)", "영업이익(A)", "매출원가(S)Tot"],
            ["S1", "테스트대리점1", "CODE_A", 100000, 10, 10000, 70000],
        ]
        buf = _write_workbook_from_grids({"DC율": dc율, "프로모션 기준": 프로모션기준, "DATA": data})
        with pytest.raises(ColumnNotFoundError, match="산식유형"):
            build_sales_like_table(buf, "2026-08")


class TestUnsupportedWorkbookFormat:
    def test_missing_dc율_sheet_raises_clear_error(self):
        buf = _write_workbook_from_grids({"다른시트": [["x"]]})
        with pytest.raises(ColumnNotFoundError, match="DC율"):
            build_sales_like_table(buf, "2026-01")
