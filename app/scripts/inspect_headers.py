"""data_sample.xlsx 시트의 병합 헤더를 복원해 컬럼별 전체 경로(상위 라벨까지 합친 튜플)를
출력한다. 실제 컬럼 매핑을 하드코딩하기 전에 검증용으로 쓴다.

실행: python app/scripts/inspect_headers.py <시트명> <헤더행수> <데이터시작행>
예:   python app/scripts/inspect_headers.py 총괄표 3 4
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
WORKBOOK_PATH = ROOT_DIR / "data_sample.xlsx"


def build_header_paths(ws, header_rows: int) -> list[tuple[str, ...]]:
    grid: list[list[str | None]] = [
        [None] * ws.max_column for _ in range(header_rows)
    ]
    for r in range(1, header_rows + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                grid[r - 1][c - 1] = str(v).replace("\n", " ").strip()

    for merged in ws.merged_cells.ranges:
        if merged.min_row > header_rows:
            continue
        top_value = grid[merged.min_row - 1][merged.min_col - 1]
        for r in range(merged.min_row, min(merged.max_row, header_rows) + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                grid[r - 1][c - 1] = top_value

    paths: list[tuple[str, ...]] = []
    for c in range(ws.max_column):
        labels: list[str] = []
        prev = None
        for r in range(header_rows):
            v = grid[r][c]
            if v and v != prev:
                labels.append(v)
                prev = v
        paths.append(tuple(labels))
    return paths


def main() -> None:
    sheet_name = sys.argv[1]
    header_rows = int(sys.argv[2])
    data_row = int(sys.argv[3]) if len(sys.argv) > 3 else header_rows + 1
    out_path = sys.argv[4] if len(sys.argv) > 4 else None

    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=False)
    ws = wb[sheet_name]
    paths = build_header_paths(ws, header_rows)
    lines = []
    for idx, path in enumerate(paths):
        sample = ws.cell(row=data_row, column=idx + 1).value
        lines.append(
            f"col{idx} ({openpyxl.utils.get_column_letter(idx + 1)}): {path!r}  sample={sample!r}"
        )
    text = "\n".join(lines)
    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(text)
    else:
        import io

        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
        print(text)


if __name__ == "__main__":
    main()
